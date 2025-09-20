"""
python linkedin_job_bot/jobbot/summarizer_generate_point.py
python -m linkedin_job_bot.jobbot.summarizer_generate_point
"""

import cohere
import re
from pathlib import Path
from docx import Document
import os
import sys

sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..')))
from jobbot.config import settings
co = cohere.Client(settings["COHERE_API_KEY3"])

PROJECTS = [
    "Project #1 – Insights Content",
    "PROJECT #2 – JAZZTRACK MANAGEMENT SYSTEM"
]

def summarize_post(job_description: str, project_tag: str) -> list[str]:
    print(f"Summarizing job for: {project_tag}")
    response = co.chat(
        message=(
            f"Summarize the following job description into very short and easy bullet points and each bullet point character length should be in between 65-85 characters"
            f"suitable for a resume under '{project_tag}'. "
            f"Start each point with a bullet (•) and use a technical CV tone.\n\n{job_description}"
        ),
        temperature=0.5,
        max_tokens=300,
    )
    raw = response.text.strip()

    cleaned: list[str] = []
    for line in raw.splitlines():
        text = line.strip()
        if not text.startswith("•"):
            continue
        text = text.replace("*", "").strip()
        text = re.sub(r"^•\s*", "", text)
        cleaned.append(text)
        if len(cleaned) == 4:
            break

    return cleaned


def extract_professional_summary(docx_path: str) -> str:
    print(f"Reading CV from: {docx_path}")
    doc = Document(docx_path)

    for i, para in enumerate(doc.paragraphs):
        if "Professional Summary" in para.text:
            print(f"Found 'Professional Summary' at index {i}")
            summary_index = i + 1
            if summary_index < len(doc.paragraphs):
                return doc.paragraphs[summary_index].text.strip()
            else:
                raise ValueError("No content found under 'Professional Summary'.")
    raise ValueError("'Professional Summary' section not found in the CV.")


def merge_summary_with_job(existing_summary: str, job_description: str) -> str:
    print("Calling Cohere to merge existing summary with job description")
    response = co.chat(
        message=(
            f"I have this existing professional summary:\n\n{existing_summary}\n\n"
            f"And this is a job summary:\n\n{job_description}\n\n"
            f"Now rewrite only the professional summary as a single clean paragraph. "
            f"Avoid any explanation or introduction. Just the summary in a technical, concise CV tone."
        ),
        temperature=0.4,
        max_tokens=300
    )
    return response.text.strip()


def generate_cv_roles(description: str, docx_path: str) -> str:
    updated_roles: dict[str, dict[str, str]] = {}

    # Summarize each project
    for project in PROJECTS:
        bullets = summarize_post(description, project)
        points_dict = { f"point{i+1}": bullet for i, bullet in enumerate(bullets) }
        updated_roles[project] = points_dict

    try: 
        existing_summary = extract_professional_summary(docx_path)
    except Exception as e:
        print(str(e))
        existing_summary = ""

    final_summary = merge_summary_with_job(existing_summary, description)
    # Clean markdown and extra spacing in the summary
    final_summary = re.sub(r"\*\*|__", "", final_summary).replace("\n", " ").strip()
    updated_roles["Summary"] = final_summary

    output_path = Path(__file__).resolve().parent.parent / "output" / "generated_roles.py"
    output_path.parent.mkdir(parents=True, exist_ok=True)

    with output_path.open("w", encoding="utf-8") as f:
        f.write("updated_roles = {\n")
        for project, pts in updated_roles.items():
            if isinstance(pts, dict):
                f.write(f'    "{project}": {{\n')
                for key, val in pts.items():
                    safe_val = val.replace('"', '\\"')
                    f.write(f'        "{key}": "{safe_val}",\n')
                f.write("    },\n")
            else:
                safe_val = pts.replace('"', '\\"')
                f.write(f'    "Summary": "{safe_val}"')
        f.write("}\n")

    print(f"Summary and bullet points saved to: {output_path}")
    return "Success: Roles and summary generated and saved."


# ----------- Usage ----------- #

if __name__ == "__main__":
    description = "We are seeking an experienced AWS API Integration Specialist to design, develop, and maintain robust integrations using AWS services and APIs. The role involves collaborating with cross-functional teams to gather requirements, implementing efficient and secure data exchanges using tools like API Gateway, Lambda (Python), Airflow, and Step Functions, and ensuring integration performance and reliability. The candidate should have deep expertise in AWS services, RESTful APIs, GraphQL, JSON/XML, scripting (Python), and security best practices. Experience with DevOps tools, version control, and CI/CD pipelines is a plus. AWS certification and knowledge of microservices and databases like RDS or DynamoDB are preferred"
    
    docx_path = "linkedin_job_bot/input/Md_Irsad_ind.docx"
    print(generate_cv_roles(description, docx_path))
