"""
cd job_profile/specific
python cv_to_docx.py --input company/<company_slug>/md_irshad_cv_2page_<company_slug>_<date>.md --jd company/<company_slug>/<company_slug>_jd_<date>.txt
Convert Markdown CV to MS Word Document (.docx) - Strict 2-Page Layout
Matches formatting from cv_to_pdf.py for visual parity.
Now includes 30-day duplicate detection via JD hash tracking.
"""
import os
import re
import argparse
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.oxml.shared import OxmlElement

# Import duplicate tracking utility
from jd_tracker import JDTracker, extract_job_slug_from_filename, extract_date_from_filename

def add_hyperlink(paragraph, text, url, font_size=9.0, color='2C5AA0', underline=True):
    """
    Add a proper Word hyperlink to a paragraph.
    This creates a clickable hyperlink in the Word document.
    """
    # Create hyperlink element
    part = paragraph.part
    r_id = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id', r_id)
    
    # Create run element
    new_run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    
    # Set color
    c = OxmlElement('w:color')
    c.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val', color)
    rPr.append(c)
    
    # Set underline
    if underline:
        u = OxmlElement('w:u')
        u.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val', 'single')
        rPr.append(u)
    
    # Set font size
    sz = OxmlElement('w:sz')
    sz.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val', str(int(font_size * 2)))
    rPr.append(sz)
    
    new_run.append(rPr)
    new_run.text = text
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    
    return hyperlink

def clean_bold_artifacts(text):
    """Remove all bold markdown and artifact tags from text."""
    # Remove ** bold markers
    text = re.sub(r'\*\*', '', text)
    # Remove artifact tags
    text = re.sub(r'\|\|/BOLD\|\|', '', text)
    text = re.sub(r'\|\|//Bold\|\|', '', text)
    text = re.sub(r'\|\|BOLD\|\|', '', text)
    return text

def add_formatted_text_to_paragraph(paragraph, text, default_size=9.0, default_color='333333', 
                                    bold=False, italic=False, force_unbold=False):
    """
    Add text with inline formatting to a Word paragraph.
    Handles **bold**, *italic*, and [links](url) properly.
    If force_unbold=True, ignores all bold markers and renders as normal text.
    """
    # If force_unbold, clean all bold markers first
    if force_unbold:
        text = clean_bold_artifacts(text)
    
    # Extract and process links using regex
    link_pattern = re.compile(r'\[([^\]]+)\]\(([^)]+)\)')
    
    # Split text by links
    parts = []
    last_end = 0
    
    for match in link_pattern.finditer(text):
        # Add text before link
        if match.start() > last_end:
            parts.append(('text', text[last_end:match.start()]))
        
        # Add link
        parts.append(('link', match.group(1), match.group(2)))
        last_end = match.end()
    
    # Add remaining text
    if last_end < len(text):
        parts.append(('text', text[last_end:]))
    
    # Process each part
    for part in parts:
        if part[0] == 'link':
            link_text, link_url = part[1], part[2]
            add_hyperlink(paragraph, link_text, link_url, font_size=default_size, color='2C5AA0')
        else:
            # Process regular text with bold/italic
            segment_text = part[1]
            
            # Parse bold and italic
            bold_pattern = re.compile(r'\*\*(.*?)\*\*')
            italic_pattern = re.compile(r'(?<!\*)\*([^*]+)\*(?!\*)')
            
            # Track position
            pos = 0
            segments = []
            
            # Find all bold/italic patterns
            for bold_match in bold_pattern.finditer(segment_text):
                if bold_match.start() > pos:
                    segments.append(('normal', segment_text[pos:bold_match.start()]))
                segments.append(('bold', bold_match.group(1)))
                pos = bold_match.end()
            
            if pos < len(segment_text):
                remaining = segment_text[pos:]
                # Check for italic in remaining
                italic_pos = 0
                for italic_match in italic_pattern.finditer(remaining):
                    if italic_match.start() > italic_pos:
                        segments.append(('normal', remaining[italic_pos:italic_match.start()]))
                    segments.append(('italic', italic_match.group(1)))
                    italic_pos = italic_match.end()
                if italic_pos < len(remaining):
                    segments.append(('normal', remaining[italic_pos:]))
            
            # If no formatting found, add as normal
            if not segments:
                segments.append(('normal', segment_text))
            
            # Add runs to paragraph
            for seg_type, seg_text in segments:
                if seg_text.strip():
                    run = paragraph.add_run(seg_text)
                    run.font.size = Pt(default_size)
                    run.font.color.rgb = RGBColor.from_string(default_color)
                    
                    if force_unbold:
                        run.font.bold = False
                    elif seg_type == 'bold':
                        run.font.bold = True
                    elif seg_type == 'italic':
                        run.font.italic = True
                    else:
                        run.font.bold = bold
                        run.font.italic = italic

def parse_markdown_to_docx(md_file, docx_file):
    """Convert markdown CV to Word document with strict 2-page layout matching PDF formatting"""
    
    with open(md_file, 'r', encoding='utf-8') as f:
        content = f.read()
    
    # Strip out HTML comment blocks
    content = re.sub(r'<!--.*?-->', '', content, flags=re.DOTALL)
    
    # Create Word document
    doc = Document()
    
    # Set page margins: Top/Bottom 0.35", Left/Right 0.4" (matching PDF)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.35)
        section.bottom_margin = Inches(0.35)
        section.left_margin = Inches(0.4)
        section.right_margin = Inches(0.4)
    
    # Set default font to Arial
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(9.0)
    
    lines = content.split('\n')
    i = 0
    
    # Track current section for context-aware formatting
    current_section = None
    current_company = None
    
    # Generic regex pattern matching project header lines like "**PROJECT_NAME – Description**"
    PROJECT_PATTERN_REGEX = re.compile(r'^\*\*([A-Z0-9_\s\/]+)\s*[–\-]\s*(.*)\*\*$', re.IGNORECASE)
    
    while i < len(lines):
        line = lines[i].strip()
        
        # Skip empty lines
        if not line:
            i += 1
            continue
        
        # Track sections
        if line.startswith('## **'):
            section_name = line[5:-2] if line.endswith('**') else line[5:]
            current_section = section_name
        
        # Track company names
        if line.startswith('### **'):
            current_company = line[6:-2] if line.endswith('**') else line[6:]
        
        # Candidate Name (# **NAME**)
        if line.startswith('# **') and line.endswith('**'):
            text = line[4:-2]
            p = doc.add_paragraph()
            p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            run = p.add_run(text)
            run.font.name = 'Arial'
            run.font.size = Pt(15.0)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
            p.paragraph_format.space_after = Pt(2)
        
        # Subtitle (Job Title)
        elif line.startswith('**') and line.endswith('**') and '|' in line and not line.startswith('### **'):
            text = line[2:-2]
            p = doc.add_paragraph()
            p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            run = p.add_run(text)
            run.font.name = 'Arial'
            run.font.size = Pt(10.0)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x2C, 0x5A, 0xA0)
            p.paragraph_format.space_after = Pt(3)
        
        # Contact Information
        elif '📞' in line or '✉️' in line or '🔗' in line or '💻' in line:
            text = line.replace('📞', '').replace('✉️', '').replace('🔗', '').replace('💻', '').strip()
            p = doc.add_paragraph()
            p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            add_formatted_text_to_paragraph(p, text, default_size=8.5, default_color='333333')
            p.paragraph_format.space_after = Pt(3)
        
        # Horizontal Rule
        elif line.startswith('---'):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
        
        # Section Heading (## **TEXT**)
        elif line.startswith('## **') and line.endswith('**'):
            text = line[5:-2]
            p = doc.add_paragraph()
            run = p.add_run(text)
            run.font.name = 'Arial'
            run.font.size = Pt(10.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(2)
        
        # Company / Role Heading (### **TEXT**)
        elif line.startswith('### **') and line.endswith('**'):
            text = line[6:-2]
            p = doc.add_paragraph()
            run = p.add_run(text)
            run.font.name = 'Arial'
            run.font.size = Pt(9.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(1)
        
        # Employment Dates (*Date Range*)
        elif line.startswith('*') and line.endswith('*') and ('–' in line or '-' in line):
            text = line[1:-1]
            p = doc.add_paragraph()
            run = p.add_run(text)
            run.font.name = 'Arial'
            run.font.size = Pt(8.5)
            run.font.italic = True
            run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            p.paragraph_format.space_after = Pt(1.5)
        
        # GENERIC PROJECT BANNERS: "**PROJECT NAME - Description**"
        elif PROJECT_PATTERN_REGEX.match(line):
            text = line[2:-2]
            p = doc.add_paragraph()
            run = p.add_run(text)
            run.font.name = 'Arial'
            run.font.size = Pt(9.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            # Apply dark blue background shading (#1A365D)
            shading_elm = parse_xml(r'<w:shd {} w:fill="1A365D"/>'.format(nsdecls('w')))
            p._element.get_or_add_pPr().append(shading_elm)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.left_indent = Inches(0.05)
        
        # Key Contributions Header
        elif line == '**Key Contributions:**':
            p = doc.add_paragraph()
            run = p.add_run('Key Contributions:')
            run.font.name = 'Arial'
            run.font.size = Pt(9.0)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            # Apply medium blue background shading (#2C5AA0)
            shading_elm = parse_xml(r'<w:shd {} w:fill="2C5AA0"/>'.format(nsdecls('w')))
            p._element.get_or_add_pPr().append(shading_elm)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.left_indent = Inches(0.05)
        
        # Bullet points (- or ✓)
        elif line.startswith('- ') or line.startswith('✓ '):
            text = line[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            
            # Check if this is under Morningstar and unbold lead-ins
            force_unbold = False
            if current_company and 'Morningstar' in current_company:
                # Check if line starts with specific patterns that should be unbolded
                morningstar_patterns = [
                    r'^\*\*AI Engineering & RAG:\*\*',
                    r'^\*\*Vector Search:\*\*',
                    r'^\*\*Data Pipelines & CDC:\*\*',
                    r'^\*\*Containerization & IaC:\*\*'
                ]
                for pattern in morningstar_patterns:
                    if re.match(pattern, text):
                        force_unbold = True
                        break
            
            add_formatted_text_to_paragraph(p, text, default_size=9.0, default_color='333333', 
                                           force_unbold=force_unbold)
            p.paragraph_format.line_spacing = Pt(11.0)
            p.paragraph_format.space_after = Pt(1.0)
            # Fix bullet indentation: set left_indent for text, first_line_indent for bullet symbol
            p.paragraph_format.left_indent = Inches(0.25)  # Text position
            p.paragraph_format.first_line_indent = Inches(-0.15)  # Bullet symbol offset (0.25 - 0.15 = 0.10")
        
        # Regular Body Text & Context-aware unbolding
        else:
            p = doc.add_paragraph()
            
            # Determine if we should force unbold based on context
            force_unbold = False
            
            # Professional Summary section - unbold everything except section header
            if current_section == 'PROFESSIONAL SUMMARY':
                force_unbold = True
            
            # Technical Skills section - unbold category headers
            elif current_section == 'TECHNICAL SKILLS':
                force_unbold = True
            
            # Tech Stack lines - unbold
            elif line.startswith('**Tech Stack:'):
                force_unbold = True
            
            # Education section - unbold
            elif current_section == 'EDUCATION & CERTIFICATIONS':
                force_unbold = True
            
            add_formatted_text_to_paragraph(p, line, default_size=9.0, default_color='333333',
                                           force_unbold=force_unbold)
            p.paragraph_format.line_spacing = Pt(11.0)
            p.paragraph_format.space_after = Pt(1.5)
        
        i += 1
    
    doc.save(docx_file)
    print(f"[SUCCESS] Generated 2-Page Word Document: {docx_file}")

if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description='Convert Markdown CV to DOCX with duplicate JD tracking',
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Examples:
  # Basic usage (no tracking):
  python cv_to_docx.py
  
  # With JD tracking:
  python cv_to_docx.py --input md_irshad_cv_2page_idc_aba_26_jul_2026.md --jd idc_aba_jd_26_jul_2026.txt
  
  # Force generation even if duplicate detected:
  python cv_to_docx.py --input md_irshad_cv_2page_idc_aba_26_jul_2026.md --jd idc_aba_jd_26_jul_2026.txt --force
        """
    )
    
    parser.add_argument(
        '--input', '-i',
        type=str,
        help='Input Markdown CV file (default: md_irshad_cv_2page_company_x.md)'
    )
    
    parser.add_argument(
        '--jd',
        type=str,
        help='Job Description file for duplicate tracking (optional)'
    )
    
    parser.add_argument(
        '--force', '-f',
        action='store_true',
        help='Force generation even if duplicate JD detected'
    )
    
    args = parser.parse_args()
    
    script_dir = os.path.dirname(os.path.abspath(__file__))
    
    # Determine input file
    if args.input:
        input_md = args.input if os.path.isabs(args.input) else os.path.join(script_dir, args.input)
    else:
        input_md = os.path.join(script_dir, "md_irshad_cv_2page_company_x.md")
    
    # Check if input file exists
    if not os.path.exists(input_md):
        print(f"[ERROR] Input file not found: {input_md}")
        exit(1)
    
    # Determine output filename (save in same directory as input file)
    input_basename = os.path.splitext(os.path.basename(input_md))[0]
    input_dir = os.path.dirname(input_md)
    output_docx = os.path.join(input_dir, f"{input_basename}.docx")
    
    print("\n" + "="*70)
    print("  CV to DOCX Converter with Duplicate JD Tracking")
    print("="*70)
    
    # JD Tracking Logic
    if args.jd:
        jd_file = args.jd if os.path.isabs(args.jd) else os.path.join(script_dir, args.jd)
        
        if not os.path.exists(jd_file):
            print(f"[WARNING] JD file not found: {jd_file}")
            print("[INFO] Proceeding without duplicate tracking")
        else:
            # Initialize tracker
            tracker = JDTracker(history_file=os.path.join(script_dir, "jd_history.json"))
            
            # Extract job slug for tracking
            job_slug = extract_job_slug_from_filename(jd_file)
            
            # Expected output files (in same directory as input)
            output_pdf = os.path.join(input_dir, f"{input_basename}.pdf")
            generated_files = [
                os.path.basename(output_pdf),
                os.path.basename(output_docx)
            ]
            
            # Check for duplicates
            should_proceed, warning = tracker.process_jd(
                jd_file_path=jd_file,
                job_slug=job_slug,
                generated_files=generated_files,
                force=args.force
            )
            
            if not should_proceed:
                # Safe console output for Windows (handles Unicode characters)
                safe_warning = warning.encode('ascii', 'replace').decode('ascii')
                print(safe_warning)
                print("\n[ABORTED] DOCX generation cancelled due to duplicate JD.")
                print("Use --force flag to override this check.")
                exit(0)
    
    # Generate Word document
    print(f"\n[INPUT]  Markdown: {os.path.basename(input_md)}")
    print(f"[OUTPUT] DOCX: {os.path.basename(output_docx)}")
    print("\n" + "-"*70)
    
    parse_markdown_to_docx(input_md, output_docx)
    
    print("-"*70)
    print(f"\n[OK] DOCX generation complete!")
    print(f"  Location: {output_docx}")
    print(f"\nTip: Run cv_to_pdf.py to generate the PDF version.")
    print("\n" + "="*70 + "\n")
